"""Offline smoke tests for file-assisted questions (no Gemini requests)."""
import json
import os
import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pptx import Presentation

from web.app import app, hour_requests, minute_requests
from web.attachments import AttachmentError, to_gemini_part


class AttachmentParserTests(unittest.TestCase):
    def test_utf8_text(self):
        part, truncated = to_gemini_part("notes.md", "今天学习 Python 列表".encode("utf-8"))
        self.assertIn("Python 列表", part.text)
        self.assertFalse(truncated)

    def test_reject_unsupported_type(self):
        with self.assertRaises(AttachmentError):
            to_gemini_part("script.exe", b"not-an-executable")

    def test_reject_fake_png(self):
        with self.assertRaises(AttachmentError):
            to_gemini_part("fake.png", b"not-an-image")

    def test_docx_text(self):
        document = Document()
        document.add_paragraph("测试 Word 文档内容")
        buffer = BytesIO()
        document.save(buffer)
        part, _ = to_gemini_part("lesson.docx", buffer.getvalue())
        self.assertIn("Word 文档", part.text)

    def test_pptx_text(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "课程幻灯片"
        buffer = BytesIO()
        presentation.save(buffer)
        part, _ = to_gemini_part("lesson.pptx", buffer.getvalue())
        self.assertIn("课程幻灯片", part.text)

    def test_xlsx_text(self):
        workbook = Workbook()
        workbook.active["A1"] = "学习进度"
        buffer = BytesIO()
        workbook.save(buffer)
        part, _ = to_gemini_part("progress.xlsx", buffer.getvalue())
        self.assertIn("学习进度", part.text)

    def test_pdf_page_cap(self):
        writer = PdfWriter()
        for _ in range(31):
            writer.add_blank_page(width=200, height=200)
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaises(AttachmentError):
            to_gemini_part("too-long.pdf", buffer.getvalue())


class APIUploadTests(unittest.TestCase):
    def setUp(self):
        os.environ["GEMINI_API_KEY"] = "placeholder-test-key"
        os.environ["GEMINI_MODEL"] = "test-model"
        minute_requests.clear()
        hour_requests.clear()
        self.client = TestClient(app)

    @patch("web.app.genai.Client")
    def test_upload_text_and_question(self, client_factory):
        client_instance = client_factory.return_value.__enter__.return_value
        client_instance.models.generate_content.return_value.text = "已读取笔记"

        response = self.client.post(
            "/api/ask-with-files",
            data={"question": "请总结附件", "mode": "quick", "history": "[]"},
            files=[("files", ("notes.txt", "学习了变量".encode("utf-8"), "text/plain"))],
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"answer": "已读取笔记", "ok": True})
        client_instance.models.generate_content.assert_called_once()

    def test_reject_bad_attachment(self):
        response = self.client.post(
            "/api/ask-with-files",
            data={"question": "这是什么", "mode": "quick", "history": "[]"},
            files=[("files", ("bad.exe", b"not allowed", "application/octet-stream"))],
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["ok"])

    @patch("web.app.genai.Client")
    def test_tutor_upload_history(self, client_factory):
        client_instance = client_factory.return_value.__enter__.return_value
        client_instance.models.generate_content.return_value.text = "继续提问"
        history = [
            {"role": "user", "text": "列表是什么"},
            {"role": "model", "text": "列表是容器"},
        ]
        response = self.client.post(
            "/api/ask-with-files",
            data={"question": "继续解释", "mode": "tutor", "history": json.dumps(history)},
            files=[("files", ("notes.txt", b"Python lists", "text/plain"))],
        )
        self.assertTrue(response.json()["ok"])
        call = client_instance.models.generate_content.call_args
        self.assertEqual(len(call.kwargs["contents"]), 3)


if __name__ == "__main__":
    unittest.main()
