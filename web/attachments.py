"""Temporary attachment parsing for the public demo.

Attachments are read in memory for one request; this module does not store them.
Office formats are converted to bounded text excerpts before being sent to Gemini.
"""
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from google.genai import types

MAX_FILES = 2
MAX_FILE_BYTES = 6 * 1024 * 1024
MAX_TOTAL_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 24000
MAX_OFFICE_UNCOMPRESSED_BYTES = 25 * 1024 * 1024

IMAGE_MIMES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py"}
OFFICE_EXTENSIONS = {".docx", ".pptx", ".xlsx"}
SUPPORTED_EXTENSIONS = set(IMAGE_MIMES) | TEXT_EXTENSIONS | OFFICE_EXTENSIONS | {".pdf"}


class AttachmentError(ValueError):
    """Unsupported, invalid or oversized attachment."""


def _check_office_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 500:
                raise AttachmentError("Office 文件内容过于复杂，请简化后重试。")
            if sum(item.file_size for item in infos) > MAX_OFFICE_UNCOMPRESSED_BYTES:
                raise AttachmentError("Office 文件解压后过大，请换用更小的文件。")
    except BadZipFile as exc:
        raise AttachmentError("Office 文件已损坏或不是有效的文件格式。") from exc


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    lines = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    slides = Presentation(BytesIO(data)).slides
    lines = []
    for slide_number, slide in enumerate(slides, 1):
        if slide_number > 40:
            lines.append("[仅提取前 40 页幻灯片]")
            break
        lines.append(f"第 {slide_number} 页")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                lines.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets[:3]:
            lines.append(f"工作表：{sheet.title}")
            for row_number, row in enumerate(
                sheet.iter_rows(max_row=150, max_col=12, values_only=True), 1
            ):
                cells = [str(value)[:200] if value is not None else "" for value in row]
                lines.append(" | ".join(cells))
            if sheet.max_row and sheet.max_row > 150:
                lines.append("[此工作表仅提取前 150 行]")
        if len(workbook.worksheets) > 3:
            lines.append("[仅提取前 3 个工作表]")
        return "\n".join(lines)
    finally:
        workbook.close()


def to_gemini_part(filename: str, data: bytes):
    """Return a Gemini Part, and whether its text was truncated."""
    safe_name = Path(filename or "未命名文件").name[:120]
    extension = Path(safe_name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise AttachmentError(
            "不支持此文件格式。请使用 PNG、JPG、WEBP、PDF、TXT、MD、CSV、JSON、PY、DOCX、PPTX 或 XLSX。"
        )

    if not data:
        raise AttachmentError(f"文件「{safe_name}」为空。")
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(f"文件「{safe_name}」超过 6 MB，请压缩后重试。")

    if extension in IMAGE_MIMES:
        valid = (
            extension in {".jpg", ".jpeg"} and data.startswith(b"\xff\xd8\xff")
            or extension == ".png" and data.startswith(b"\x89PNG\r\n\x1a\n")
            or extension == ".webp" and data[:4] == b"RIFF" and data[8:12] == b"WEBP"
        )
        if not valid:
            raise AttachmentError(f"图片「{safe_name}」格式不正确。")
        return types.Part.from_bytes(data=data, mime_type=IMAGE_MIMES[extension]), False

    if extension == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF-"):
            raise AttachmentError(f"PDF「{safe_name}」格式不正确。")
        return types.Part.from_bytes(data=data, mime_type="application/pdf"), False

    if extension in OFFICE_EXTENSIONS:
        _check_office_archive(data)
        try:
            if extension == ".docx":
                text = _extract_docx(data)
            elif extension == ".pptx":
                text = _extract_pptx(data)
            else:
                text = _extract_xlsx(data)
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError(f"无法读取「{safe_name}」，请确认文件未损坏。") from exc
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise AttachmentError(f"「{safe_name}」不是 UTF-8 文本，请转换编码后重试。") from exc

    if not text.strip():
        raise AttachmentError(f"没有从「{safe_name}」提取到可分析的文字。")
    truncated = len(text) > MAX_EXTRACTED_CHARS
    excerpt = text[:MAX_EXTRACTED_CHARS]
    part = types.Part.from_text(
        text=f"\n[附件资料开始：{safe_name}]\n{excerpt}\n[附件资料结束]\n"
    )
    return part, truncated
